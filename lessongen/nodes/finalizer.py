import json
import logging
from pathlib import Path
from typing import Any, Dict, List
import docx
from docx.shared import Pt, Inches, RGBColor
from lessongen.config import settings
from lessongen.state import LessonState
from lessongen.nodes.models import AttemptLog

logger = logging.getLogger("lessongen.finalizer")

def finalize_lesson_outputs(state: LessonState) -> Dict[str, Any]:
    """Writes all output files: final_lesson.md, final_lesson.docx, rejection_log (.md/.json), run_trace.json."""
    output_dir = settings.OUTPUTS_DIR
    output_dir.mkdir(parents=True, exist_ok=True)

    final_status = state.get("final_status", "PASSED")
    draft_lesson = state.get("draft_lesson", "")
    rejection_log: List[AttemptLog] = state.get("rejection_log", [])
    trace: List[Dict[str, Any]] = state.get("trace", [])

    # 1. Format Markdown Lesson
    md_content = draft_lesson
    if final_status == "NEEDS_HUMAN_REVIEW":
        header = (
            "> [!WARNING]\n"
            "> **STATUS: NEEDS HUMAN REVIEW**\n"
            "> This lesson failed automated evaluation checkpoints after maximum retries.\n"
            "> A human editor must inspect and resolve the logged failures before publishing.\n\n"
        )
        md_content = header + md_content

    md_path = output_dir / "final_lesson.md"
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_content)

    # 2. Format DOCX Document using python-docx
    docx_path = output_dir / "final_lesson.docx"
    export_to_docx(draft_lesson, final_status, docx_path)

    # 3. Write Rejection Log (.json and .md)
    json_path = output_dir / "rejection_log.json"
    log_dicts = [log.model_dump() for log in rejection_log]
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(log_dicts, f, indent=2)

    md_log_path = output_dir / "rejection_log.md"
    write_rejection_log_markdown(rejection_log, final_status, md_log_path)

    # 3b. If deliberate error was injected, write demo_rejection_log.md
    if state.get("injected_error"):
        demo_log_path = output_dir / "demo_rejection_log.md"
        write_demo_rejection_log_markdown(
            rejection_log=rejection_log,
            injected_error=state.get("injected_error"),
            dest_path=demo_log_path
        )

    # 4. Write Run Trace (.json)
    trace_path = output_dir / "run_trace.json"
    with open(trace_path, "w", encoding="utf-8") as f:
        json.dump({
            "topic": state.get("topic"),
            "final_status": final_status,
            "attempts_used": state.get("current_attempt", 1),
            "trace_events": trace
        }, f, indent=2)

    logger.info(f"Final outputs successfully generated in {output_dir}")
    return {
        "final_lesson_md": str(md_path),
        "final_lesson_docx": str(docx_path),
        "rejection_log_md": str(md_log_path),
        "rejection_log_json": str(json_path),
        "run_trace_json": str(trace_path)
    }

def export_to_docx(markdown_text: str, status: str, dest_path: Path) -> None:
    """Exports markdown lesson to a clean, professionally formatted .docx document."""
    doc = docx.Document()

    # Document Title / Header
    if status == "NEEDS_HUMAN_REVIEW":
        warning_p = doc.add_paragraph()
        warning_run = warning_p.add_run("STATUS: NEEDS HUMAN REVIEW (Failed Automated Quality Gates)")
        warning_run.bold = True
        warning_run.font.color.rgb = RGBColor(180, 40, 40)
        warning_run.font.size = Pt(13)

    for line in markdown_text.splitlines():
        trimmed = line.strip()
        if not trimmed:
            continue
        if trimmed.startswith("# "):
            p = doc.add_heading(trimmed[2:], level=1)
            p.runs[0].font.size = Pt(20)
            p.runs[0].font.color.rgb = RGBColor(20, 40, 80)
        elif trimmed.startswith("## "):
            p = doc.add_heading(trimmed[3:], level=2)
            p.runs[0].font.size = Pt(14)
            p.runs[0].font.color.rgb = RGBColor(30, 60, 110)
        elif trimmed.startswith("### "):
            p = doc.add_heading(trimmed[4:], level=3)
            p.runs[0].font.size = Pt(12)
        elif trimmed.startswith("- ") or trimmed.startswith("* "):
            doc.add_paragraph(trimmed[2:], style="List Bullet")
        elif trimmed.startswith("1. ") or trimmed.startswith("2. ") or trimmed.startswith("3. ") or trimmed.startswith("4. ") or trimmed.startswith("5. "):
            doc.add_paragraph(trimmed[3:], style="List Number")
        elif trimmed == "---":
            continue
        else:
            doc.add_paragraph(trimmed)

    doc.save(str(dest_path))

def write_rejection_log_markdown(logs: List[AttemptLog], final_status: str, dest_path: Path) -> None:
    """Writes human-readable rejection log markdown."""
    with open(dest_path, "w", encoding="utf-8") as f:
        f.write("# Lesson Quality Evaluation & Rejection Log\n\n")
        f.write(f"**Final System Verdict:** `{final_status}`\n")
        f.write(f"**Total Generation Attempts:** {len(logs)}\n\n")
        f.write("---\n\n")

        for log in logs:
            status_badge = "PASSED" if log.passed else "FAILED / REJECTED"
            f.write(f"## Attempt {log.attempt} — Verdict: {status_badge}\n\n")
            
            if log.metrics:
                f.write("### Quantitative Metrics:\n")
                for k, v in log.metrics.items():
                    f.write(f"- **{k.replace('_', ' ').title()}**: {v}\n")
                f.write("\n")

            if log.passed:
                f.write("> **Success:** All 16 hard pass/fail checkpoints satisfied. Ready for publication.\n\n")
            else:
                f.write(f"### Failed Checkpoints ({len(log.failed_checkpoints)}):\n\n")
                for cp_id in log.failed_checkpoints:
                    reason = log.reasons.get(cp_id, "Pass condition not met.")
                    evidence = log.evidence.get(cp_id, "No direct quote available.")
                    f.write(f"#### Checkpoint `{cp_id}`\n")
                    f.write(f"- **Failure Reason**: {reason}\n")
                    f.write(f"- **Direct Evidence Cited**: *\"{evidence}\"*\n\n")

                if log.fix_instruction:
                    f.write("### Reflector Fix Instructions:\n")
                    f.write(f"```text\n{log.fix_instruction}\n```\n\n")

                if log.diff_summary:
                    f.write("### Targeted Diff Summary:\n")
                    f.write(f"{log.diff_summary}\n\n")

            f.write("---\n\n")

def write_demo_rejection_log_markdown(rejection_log: List[AttemptLog], injected_error: str, dest_path: Path) -> None:
    """Writes specialized rejection log markdown highlighting the deliberate error demonstration for video."""
    with open(dest_path, "w", encoding="utf-8") as f:
        f.write("# Deliberate Error Demonstration Log\n\n")
        f.write(f"**Demonstration Scenario:** Injected `{injected_error}` error into candidate lesson.\n")
        f.write(f"**Objective:** Verify that Evaluator Layer catches the seeded error, cites direct quote evidence, and Reflector heals the draft on retry.\n\n")
        f.write("---\n\n")

        for log in rejection_log:
            status_badge = "PASSED (Error Fixed)" if log.passed else f"FAILED (Seeded Error Caught)"
            f.write(f"## Attempt {log.attempt} — Verdict: {status_badge}\n\n")
            if log.passed:
                f.write("> **Recovery Confirmed:** The candidate lesson was regenerated, removing all seeded errors. All 16 checkpoints passed.\n\n")
            else:
                f.write(f"### Caught Checkpoints ({len(log.failed_checkpoints)}):\n\n")
                for cp_id in log.failed_checkpoints:
                    reason = log.reasons.get(cp_id, "Pass condition violated.")
                    evidence = log.evidence.get(cp_id, "No quote cited.")
                    f.write(f"#### Checkpoint `{cp_id}`\n")
                    f.write(f"- **Evaluator Diagnosis**: {reason}\n")
                    f.write(f"- **Direct Evidence Caught**: *\"{evidence}\"*\n\n")

                if log.fix_instruction:
                    f.write("### Reflector Action Plan:\n")
                    f.write(f"```text\n{log.fix_instruction}\n```\n\n")

                if log.diff_summary:
                    f.write(f"**Diff Summary:** {log.diff_summary}\n\n")
            f.write("---\n\n")

